import tkinter as tk
from tkinter import ttk, messagebox
from tkcalendar import DateEntry
from datetime import timedelta
from finance_app.models.transaction import Transaction
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
import matplotlib.pyplot as plt
from datetime import datetime
from tkinter import filedialog
from fpdf import FPDF
import sys
from finance_app.importers.pko_parser import parse_pko_pdf


def on_tab_change(self, event):
    tab = self.notebook.index(self.notebook.select())
    self.current_type = "expense" if tab == 0 else "income"
    self.load_categories()
    self.refresh_view()
    self.load_accounts()


def set_filter(self, value):
    self.date_filter = value
    self.refresh_view()


def get_date_range(self):
    today = datetime.today()
    if self.date_filter == "day":
        return today, today
    elif self.date_filter == "week":
        return today - timedelta(days=7), today
    elif self.date_filter == "month":
        return today.replace(day=1), today
    elif self.date_filter == "custom":
        start = datetime.strptime(self.date_from.get(), "%m/%d/%y")
        end = datetime.strptime(self.date_to.get(), "%m/%d/%y")
        return start, end


def load_categories(self):
    categories = self.cat_repo.get_by_type(self.current_type)
    self.cat_combo["values"] = categories
    if categories:
        self.cat_combo.current(0)


def refresh_view(self):
    start_date, end_date = self.get_date_range()
    account_name = self.filter_account_var.get()
    account_id = self.account_map.get(account_name) if account_name != "All accounts" else None
    category = self.cat_var.get() if self.filter_by_cat_var.get() else None

    transactions = self.repo.get_by_type_and_period(
        self.current_type, start_date, end_date, category, account_id
    )

    total = sum(tx.amount for tx in transactions)

    self.total_label.config(text=f"Amount: {total:.2f}")

    for row in self.table.get_children():
        self.table.delete(row)
    for tx in transactions:
        account = self.account_repo.get_by_id(tx.account_id)
        account_name = account.name if account else "—"
        self.table.insert("", "end", values=(tx.date, tx.amount, tx.category, tx.description, account_name))


def add_new_category(self):
    win = tk.Toplevel(self.root)
    win.title("New category")

    tk.Label(win, text="Category name:").pack(pady=5)
    name_entry = tk.Entry(win)
    name_entry.pack(padx=10)

    def save_category():
        name = name_entry.get().strip()
        if name:
            self.cat_repo.add_category(name, self.current_type)
            self.load_categories()
            self.cat_var.set(name)
            win.destroy()

    tk.Button(win, text="Save", command=save_category).pack(pady=5)


def add_transaction(self):
    try:
        amount = float(self.amount_var.get())
        category = self.cat_var.get()
        description = self.desc_var.get()
        date = self.date_var.get_date().strftime("%Y-%m-%d")
        account_name = self.account_var.get()

        if not category:
            messagebox.showwarning("Category", "Choose the category!")
            return
        if not account_name:
            messagebox.showwarning("Account", "Choose the account!")
            return

        account_id = self.account_map.get(account_name)
        if not account_id:
            messagebox.showerror("Error", "Such account is not found.")
            return

        tx = Transaction(
            id=None,
            date=date,
            amount=amount,
            category=category,
            description=description,
            type=self.current_type,
            account_id=account_id
        )

        self.repo.add(tx)

        account = self.account_repo.get_by_id(account_id)
        if self.current_type == "expense":
            account.balance -= amount
        else:
            account.balance += amount
        self.account_repo.update_account(account_id, account.name, account.balance)

        self.refresh_view()
        self.refresh_accounts_table()

        self.amount_var.set("")
        self.desc_var.set("")

    except ValueError:
        messagebox.showerror("Error", "Enter the correct amount.")


def show_report(self):
    start_date, end_date = self.get_date_range()
    category_filter = self.cat_var.get() if self.filter_by_cat_var.get() else None
    account_name = self.filter_account_var.get()
    account_id = self.account_map.get(account_name) if account_name != "All accounts" else None

    transactions = self.repo.get_by_type_and_period(
        self.current_type, start_date, end_date, category_filter, account_id
    )

    if not transactions:
        messagebox.showinfo("Report", "There are no data for the report.")
        return

    data = {}
    for tx in transactions:
        data[tx.category] = data.get(tx.category, 0) + tx.amount

    categories = list(data.keys())
    values = list(data.values())

    report_window = tk.Toplevel(self.root)
    report_window.title("Report by categories")
    report_window.geometry("600x450")

    fig, ax = plt.subplots(figsize=(6, 4), dpi=100)
    ax.pie(values, labels=categories, autopct="%1.1f%%", startangle=140)
    ax.axis("equal")
    title = "Expenses by categories" if self.current_type == "expense" else "Income by categories"
    ax.set_title(f"{title}\n{start_date.strftime('%Y-%m-%d')} – {end_date.strftime('%Y-%m-%d')}")

    canvas = FigureCanvasTkAgg(fig, master=report_window)
    canvas.draw()
    canvas.get_tk_widget().pack(fill="both", expand=True, padx=10, pady=10)


def delete_transaction(self):
    selected = self.table.selection()
    if not selected:
        messagebox.showwarning("Deleting", "Select the transaction.")
        return

    item = self.table.item(selected[0])
    values = item["values"]
    date, amount, category, description, _ = values

    if not messagebox.askyesno("Confirmation", f"Delete the transaction with the amount of {amount}?"):
        return

    self.repo.delete_transaction(date, float(amount), category, description)
    self.refresh_view()


def edit_transaction(self):
    selected = self.table.selection()
    if not selected:
        messagebox.showwarning("Editing", "Select the transaction.")
        return

    item = self.table.item(selected[0])
    old_date, old_amount, old_category, old_description, _ = item["values"]

    win = tk.Toplevel(self.root)
    win.title("Edit the transaction")

    tk.Label(win, text="Amount:").grid(row=0, column=0)
    amount_var = tk.StringVar(value=str(old_amount))
    tk.Entry(win, textvariable=amount_var).grid(row=0, column=1)

    tk.Label(win, text="Category:").grid(row=1, column=0)
    cat_var = tk.StringVar(value=old_category)
    cat_combo = ttk.Combobox(win, textvariable=cat_var,
                             values=self.cat_repo.get_by_type(self.current_type),
                             state="readonly")
    cat_combo.grid(row=1, column=1)

    tk.Label(win, text="Description:").grid(row=2, column=0)
    desc_var = tk.StringVar(value=old_description)
    tk.Entry(win, textvariable=desc_var).grid(row=2, column=1)

    tk.Label(win, text="Date:").grid(row=3, column=0)
    date_var = DateEntry(win, width=10)
    parsed_date = datetime.strptime(old_date, "%Y-%m-%d").date()
    date_var.set_date(parsed_date)
    date_var.grid(row=3, column=1)

    def save_changes():
        try:
            new_tx = Transaction(
                id=None,
                date=date_var.get_date().strftime("%Y-%m-%d"),
                amount=float(amount_var.get()),
                category=cat_var.get(),
                description=desc_var.get(),
                type=self.current_type,
                account_id=None
            )
            self.repo.update_transaction(
                old_date, float(old_amount), old_category, old_description, new_tx
            )
            self.refresh_view()
            win.destroy()
        except ValueError:
            messagebox.showerror("Error", "Incorrect data.")

    tk.Button(win, text="Save", command=save_changes).grid(row=4, column=0, columnspan=2, pady=10)



def import_pdf(self):
    from tkinter import filedialog

    filepath = filedialog.askopenfilename(filetypes=[("PDF files", "*.pdf")])
    if not filepath:
        return

    transactions = parse_pko_pdf(filepath)
    for tx in transactions:
        self.repo.add(tx)

    messagebox.showinfo("Import completed", f"Imported {len(transactions)} transactions.")
    self.refresh_view()


def export_as_csv(self):
    from tkinter import filedialog
    import csv

    start_date, end_date = self.get_date_range()
    transactions = self.repo.get_by_type_and_period(self.current_type, start_date, end_date)

    if not transactions:
        messagebox.showinfo("Export", "There no transactions to export.")
        return

    filepath = filedialog.asksaveasfilename(
        defaultextension=".csv",
        filetypes=[("CSV files", "*.csv")],
        title="Save as CSV"
    )

    if not filepath:
        return

    with open(filepath, mode='w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(["Date", "Amount", "Category", "Description", "Type"])
        for tx in transactions:
            writer.writerow([tx.date, tx.amount, tx.category, tx.description, tx.type])

    messagebox.showinfo("Export completed", f"Exported {len(transactions)} transactions as CSV.")


def export_as_pdf(self):
    import os

    start_date, end_date = self.get_date_range()
    transactions = self.repo.get_by_type_and_period(self.current_type, start_date, end_date)

    if not transactions:
        messagebox.showinfo("Export", "There are no transactions to export.")
        return

    filepath = filedialog.asksaveasfilename(
        defaultextension=".pdf",
        filetypes=[("PDF files", "*.pdf")],
        title="Save as PDF"
    )

    if not filepath:
        return

    pdf = FPDF()
    pdf.add_page()

    if hasattr(sys, '_MEIPASS'):
        base_dir = os.path.join(sys._MEIPASS, "fonts")
    else:
        base_dir = os.path.join(os.path.dirname(__file__), "..", "fonts")

    font_path = os.path.join(base_dir, "DejaVuSans.ttf")
    font_path = os.path.abspath(font_path)  # полный путь
    pdf.add_font('DejaVu', '', font_path, uni=True)
    pdf.set_font('DejaVu', size=10)

    pdf.cell(200, 10, txt="Transactions report", ln=True, align='C')
    pdf.cell(200, 10,
             txt=f"{self.current_type.upper()} с {start_date.strftime('%Y-%m-%d')} по {end_date.strftime('%Y-%m-%d')}",
             ln=True, align='C')
    pdf.ln(10)

    pdf.set_font('DejaVu', size=9)
    pdf.cell(30, 10, "Date", 1)
    pdf.cell(25, 10, "Amount", 1)
    pdf.cell(40, 10, "Category", 1)
    pdf.cell(95, 10, "Description", 1)
    pdf.ln()

    for tx in transactions:
        pdf.cell(30, 10, tx.date, 1)
        pdf.cell(25, 10, f"{tx.amount:.2f}", 1)
        pdf.cell(40, 10, tx.category[:15], 1)
        pdf.cell(95, 10, tx.description[:35], 1)
        pdf.ln()

    pdf.output(filepath)
    messagebox.showinfo("Export completed", f"Exported {len(transactions)} transactions as PDF.")


def export_as_excel(self):
    from tkinter import filedialog
    from openpyxl import Workbook

    start_date, end_date = self.get_date_range()
    transactions = self.repo.get_by_type_and_period(self.current_type, start_date, end_date)

    if not transactions:
        messagebox.showinfo("Export", "There are no transactions to export.")
        return

    filepath = filedialog.asksaveasfilename(
        defaultextension=".xlsx",
        filetypes=[("Excel files", "*.xlsx")],
        title="Save as Excel"
    )

    if not filepath:
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"

    ws.append(["Date", "Amount", "Category", "Description", "Type"])

    for tx in transactions:
        ws.append([tx.date, tx.amount, tx.category, tx.description, tx.type])

    wb.save(filepath)
    messagebox.showinfo("Export completed", f"Exported {len(transactions)} transactions as Excel.")


def setup_accounts_tab(self):
    self.accounts_frame = ttk.Frame(self.notebook)
    self.notebook.add(self.accounts_frame, text="Accounts")

    self.accounts_table = ttk.Treeview(
        self.accounts_frame,
        columns=("Name", "Balance"),
        show="headings"
    )
    self.accounts_table.heading("Name", text="Name")
    self.accounts_table.heading("Balance", text="Balance")
    self.accounts_table.pack(fill="both", expand=True, pady=10, padx=10)

    actions = ttk.Frame(self.accounts_frame)
    actions.pack(pady=5)

    ttk.Button(actions, text="+ Add account", command=self.add_account_window).pack(side="left", padx=5)
    ttk.Button(actions, text="Edit", command=self.edit_account_window).pack(side="left", padx=5)
    ttk.Button(actions, text="Delete", command=self.delete_account).pack(side="left", padx=5)

    self.refresh_accounts_table()


def refresh_accounts_table(self):
    for row in self.accounts_table.get_children():
        self.accounts_table.delete(row)
    for acc in self.account_repo.get_all():
        self.accounts_table.insert("", "end", values=(acc.name, f"{acc.balance:.2f}"))


def add_account_window(self):
    win = tk.Toplevel(self.root)
    win.title("Add account")

    tk.Label(win, text="Name:").grid(row=0, column=0)
    name_var = tk.StringVar()
    tk.Entry(win, textvariable=name_var).grid(row=0, column=1)

    tk.Label(win, text="Balance:").grid(row=1, column=0)
    balance_var = tk.StringVar()
    tk.Entry(win, textvariable=balance_var).grid(row=1, column=1)

    def save():
        try:
            name = name_var.get().strip()
            balance = float(balance_var.get())
            if name:
                self.account_repo.add_account(name, balance)
                self.refresh_accounts_table()
                win.destroy()
        except ValueError:
            messagebox.showerror("Error", "Incorrect balance")

    tk.Button(win, text="Save", command=save).grid(row=2, column=0, columnspan=2, pady=10)


def edit_account_window(self):
    selected = self.accounts_table.selection()
    if not selected:
        messagebox.showwarning("Edit", "Choose the account")
        return

    item = self.accounts_table.item(selected[0])
    old_name, old_balance = item["values"]
    account = self.account_repo.get_by_name(old_name)

    win = tk.Toplevel(self.root)
    win.title("Edit account")

    tk.Label(win, text="Name:").grid(row=0, column=0)
    name_var = tk.StringVar(value=old_name)
    tk.Entry(win, textvariable=name_var).grid(row=0, column=1)

    tk.Label(win, text="Balance:").grid(row=1, column=0)
    balance_var = tk.StringVar(value=old_balance)
    tk.Entry(win, textvariable=balance_var).grid(row=1, column=1)

    def save():
        try:
            new_name = name_var.get().strip()
            new_balance = float(balance_var.get())
            self.account_repo.update_account(account.id, new_name, new_balance)
            self.refresh_accounts_table()
            win.destroy()
        except ValueError:
            messagebox.showerror("Error", "Incorrect balance")

    tk.Button(win, text="Save", command=save).grid(row=2, column=0, columnspan=2, pady=10)


def delete_account(self):
    selected = self.accounts_table.selection()
    if not selected:
        messagebox.showwarning("Deleting", "Choose the account.")
        return

    item = self.accounts_table.item(selected[0])
    name = item["values"][0]
    account = self.account_repo.get_by_name(name)

    if messagebox.askyesno("Deleting", f"Delet the account '{name}'?"):
        self.account_repo.delete_account(account.id)
        self.refresh_accounts_table()


def load_accounts(self):
    accounts = self.account_repo.get_all()
    self.account_map = {acc.name: acc.id for acc in accounts}

    self.account_combo["values"] = list(self.account_map.keys())
    if accounts:
        self.account_combo.current(0)

    self.filter_account_combo["values"] = ["All the accounts"] + list(self.account_map.keys())
    self.filter_account_combo.current(0)


def on_tab_change(self, event):
    tab = self.notebook.index(self.notebook.select())

    if tab == 1:
        self.current_type = "expense"
    elif tab == 2:
        self.current_type = "income"
    elif tab == 0:
        return
    else:
        return

    self.load_categories()
    self.refresh_view()
    self.load_accounts()


def sort_by_column(self, col):
    col_map = {
        "Date": 0,
        "Amount": 1,
        "Category": 2,
        "Description": 3
    }

    col_index = col_map[col]
    data = [(self.table.set(k, col), k) for k in self.table.get_children()]

    if col == "Amount":
        data.sort(key=lambda t: float(t[0].replace(",", ".")),
                  reverse=self.sort_column == col and not self.sort_reverse)
    elif col == "Date":
        data.sort(key=lambda t: datetime.strptime(t[0], "%Y-%m-%d"),
                  reverse=self.sort_column == col and not self.sort_reverse)
    else:
        data.sort(key=lambda t: t[0], reverse=self.sort_column == col and not self.sort_reverse)

    for index, (val, k) in enumerate(data):
        self.table.move(k, "", index)

    self.sort_column = col
    self.sort_reverse = not (self.sort_column == col and self.sort_reverse)

